"""File input and deterministic customer record cleaning, independent of the API."""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator

from openpyxl import load_workbook


COLUMNS = ("name", "email", "phone", "birth_date")
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
EMAIL_LOCAL = re.compile(r"^[A-Za-z0-9!#$%&'*+/=?^_`{|}~.-]+$")
DOMAIN_LABEL = re.compile(r"^[A-Za-z0-9](?:[A-Za-z0-9-]*[A-Za-z0-9])?$")
PHONE_CHARS = re.compile(r"^[0-9+().\-\s]+$")
PARTICLES = {"de", "da", "do", "das", "dos"}


class InputError(ValueError):
    """The uploaded file cannot be processed as a customer sheet."""

    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(message)


@dataclass(frozen=True)
class ProcessResult:
    summary: dict[str, int]
    clean_csv: bytes
    errors_csv: bytes


def _text(value: Any) -> str:
    if getattr(value, "data_type", None) == "f":
        value = value.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _nonempty(values: Iterable[Any]) -> bool:
    return any(_text(value).strip() for value in values)


def _csv_rows(data: bytes) -> Iterator[tuple[int, list[Any]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InputError("INVALID_ENCODING", "CSV must use UTF-8") from exc
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    try:
        while True:
            source_row = reader.line_num + 1
            try:
                row = next(reader)
            except StopIteration:
                break
            yield source_row, row
    except csv.Error as exc:
        raise InputError("INVALID_CSV", "Malformed CSV") from exc


def _xlsx_rows(data: bytes) -> Iterator[tuple[int, list[Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise InputError("INVALID_XLSX", "Unreadable XLSX") from exc
    try:
        sheet = workbook.worksheets[0]
        for cells in sheet.iter_rows():
            values = [cell if cell.data_type == "f" else cell.value for cell in cells]
            yield cells[0].row, values
    finally:
        workbook.close()


def _header(values: list[Any]) -> dict[str, int]:
    names = [_text(value).strip() for value in values]
    if len(names) != 4 or len(set(names)) != 4 or set(names) != set(COLUMNS):
        raise InputError("INVALID_HEADER", "Header must contain exactly name, email, phone, and birth_date")
    return {name: index for index, name in enumerate(names)}


def _email(value: str) -> str | None:
    if value.count("@") != 1 or any(char.isspace() for char in value):
        return None
    local, domain = value.split("@")
    if not local or local.startswith(".") or local.endswith(".") or ".." in local:
        return None
    if not EMAIL_LOCAL.fullmatch(local):
        return None
    labels = domain.split(".")
    if len(labels) < 2 or any(not DOMAIN_LABEL.fullmatch(label) for label in labels):
        return None
    return value


def _phone(value: str) -> str | None:
    if not PHONE_CHARS.fullmatch(value) or value.count("+") > 1 or ("+" in value and not value.startswith("+55")):
        return None
    digits = re.sub(r"\D", "", value)
    if value.startswith("+55") or len(digits) in (12, 13) and digits.startswith("55"):
        if not digits.startswith("55"):
            return None
        digits = digits[2:]
    if len(digits) not in (10, 11) or not 11 <= int(digits[:2]) <= 99:
        return None
    return "+55" + digits


def _birth_date(value: Any, today: date) -> tuple[str | None, str | None]:
    parsed: date
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            return None, "INVALID_DATE"
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        raw = _text(value).strip()
        try:
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
                parsed = date.fromisoformat(raw)
            elif re.fullmatch(r"\d{2}/\d{2}/\d{4}", raw):
                parsed = datetime.strptime(raw, "%d/%m/%Y").date()
            else:
                return None, "INVALID_DATE"
        except ValueError:
            return None, "INVALID_DATE"
    if parsed > today:
        return None, "FUTURE_DATE"
    return parsed.isoformat(), None


def _write_csv(header: tuple[str, ...], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def process_file(data: bytes, filename: str, *, today: date | None = None) -> ProcessResult:
    """Clean a CSV/XLSX file and return stable CSV outputs plus row counts."""
    if not data:
        raise InputError("EMPTY_FILE", "Empty file")
    if len(data) > MAX_BYTES:
        raise InputError("FILE_TOO_LARGE", "File exceeds 5 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = _csv_rows(data)
    elif suffix == ".xlsx":
        rows = _xlsx_rows(data)
    else:
        raise InputError("UNSUPPORTED_FORMAT", "Upload a CSV or XLSX file")
    current_date = today or date.today()
    try:
        _, first = next(rows)
    except StopIteration as exc:
        raise InputError("EMPTY_FILE", "File has no header") from exc
    mapping = _header(first)
    accepted: list[list[str]] = []
    issues: list[list[Any]] = []
    first_email_row: dict[str, int] = {}
    summary = {"received": 0, "valid": 0, "invalid": 0, "duplicate": 0}

    for source_row, cells in rows:
        if not _nonempty(cells):
            continue
        summary["received"] += 1
        if summary["received"] > MAX_ROWS:
            raise InputError("TOO_MANY_ROWS", "File exceeds 10,000 data rows")
        if len(cells) != 4:
            issues.append([source_row, "row", json.dumps([_text(value) for value in cells], ensure_ascii=False), "INVALID_ROW", "Row must contain four cells"])
            summary["invalid"] += 1
            continue

        raw = {column: cells[index] for column, index in mapping.items()}
        values = {column: _text(raw[column]) for column in COLUMNS}
        line_issues: list[list[Any]] = []

        def issue(column: str, code: str, message: str) -> None:
            line_issues.append([source_row, column, values[column], code, message])

        for column in COLUMNS:
            if getattr(raw[column], "data_type", None) == "f":
                issue(column, "FORMULA_NOT_ALLOWED", "Formulas are not allowed")

        name_words = values["name"].split()
        name = " ".join(word.lower() if index > 0 and word.lower() in PARTICLES else word.capitalize() for index, word in enumerate(name_words))
        if not name and not any(item[1] == "name" for item in line_issues):
            issue("name", "REQUIRED", "Name is required")

        email = values["email"].strip().lower()
        if not email and not any(item[1] == "email" for item in line_issues):
            issue("email", "REQUIRED", "Email is required")
        elif email and _email(email) is None and not any(item[1] == "email" for item in line_issues):
            issue("email", "INVALID_EMAIL", "Invalid email")

        phone_raw = values["phone"].strip()
        phone = "" if not phone_raw else _phone(phone_raw)
        if phone is None and not any(item[1] == "phone" for item in line_issues):
            issue("phone", "INVALID_PHONE", "Phone must include an area code and 10 or 11 national digits")

        birth_raw = values["birth_date"].strip()
        birth = ""
        if birth_raw:
            birth, date_error = _birth_date(raw["birth_date"], current_date)
            if date_error and not any(item[1] == "birth_date" for item in line_issues):
                issue("birth_date", date_error, "Birth date is after the processing date" if date_error == "FUTURE_DATE" else "Invalid or nonexistent date")

        if line_issues:
            issues.extend(line_issues)
            summary["invalid"] += 1
        elif email in first_email_row:
            issues.append([source_row, "email", values["email"], "DUPLICATE_EMAIL", f"Duplicate email; first valid occurrence on row {first_email_row[email]}"])
            summary["duplicate"] += 1
        else:
            first_email_row[email] = source_row
            accepted.append([name, email, phone, birth])
            summary["valid"] += 1

    return ProcessResult(
        summary=summary,
        clean_csv=_write_csv(COLUMNS, accepted),
        errors_csv=_write_csv(("source_row", "column", "original_value", "code", "message"), issues),
    )
